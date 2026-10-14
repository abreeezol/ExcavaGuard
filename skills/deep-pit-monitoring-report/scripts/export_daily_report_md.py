# -*- coding: utf-8 -*-
"""
export_daily_report_md.py — 把 gen_daily_report_rag.py 产出的监测日报 xlsx
（监测日报 + 各附录独立 sheet + 超限测点清单）导出为两份 Markdown：

  ① 仅「监测日报」封面 sheet 的内容  ->  <stem>_封面.md
  ② 除「监测日报」外的其余 sheet      ->  <stem>_附表及清单.md

用法：
  python scripts/export_daily_report_md.py --xlsx 监测日报_附表AF_RAG.xlsx
  # 可自定义：
  #   --xlsx <日报.xlsx> --cover-sheet 监测日报
  #   --out-cover <封面.md> --out-rest <附表及清单.md>

依赖：openpyxl。
"""
from __future__ import annotations
import os, re, argparse
import openpyxl


SECTION_PAT = re.compile(r"^[一二三四五六七八九十]+、")


def _md_cell(v):
    """把单元格值转成可安全放入 Markdown 表格单元格的文本。"""
    if v is None:
        return ""
    s = str(v)
    s = s.replace("\\", "\\\\").replace("|", "\\|")
    s = s.replace("\r\n", "\n").replace("\r", "\n").replace("\n", "<br>")
    return s.strip()


def _strip_leading_bullet(s):
    s = str(s).strip()
    for pre in ("· ", "·", "•", "- ", "* "):
        if s.startswith(pre):
            return s[len(pre):].strip()
    return s


def _nonempty_count(row):
    return sum(1 for v in row if v is not None and str(v).strip() != "")


# ----------------------------- 封面 sheet 渲染 -----------------------------
def render_cover(ws):
    """把「监测日报」封面 sheet 渲染成结构化 Markdown。"""
    rows = []
    for r in range(1, ws.max_row + 1):
        a = ws.cell(r, 1).value
        b = ws.cell(r, 2).value if ws.max_column >= 2 else None
        rows.append((a, b))

    out = []
    table_buf = []

    def flush_table():
        if not table_buf:
            return
        first = table_buf[0]
        if first[0] in ("内容", "检索主题", "项目"):
            header, body = first, table_buf[1:]
        else:
            header, body = ("项目", "内容"), table_buf[:]
        out.append("| %s | %s |" % (_md_cell(header[0]), _md_cell(header[1])))
        out.append("| --- | --- |")
        for a, b in body:
            out.append("| %s | %s |" % (_md_cell(a), _md_cell(b)))
        table_buf.clear()

    for a, b in rows:
        a_s = None if a is None else str(a).strip()
        b_s = None if b is None else str(b).strip()
        a_has = bool(a_s)
        b_has = b_s is not None and b_s != ""

        if a_has and not b_has:
            flush_table()
            if SECTION_PAT.match(a_s):
                out.append("")
                out.append("## " + a_s)
            else:
                out.append("")
                out.append("# " + a_s)
        elif a_has and b_has:
            if a_s in ("·", "□"):
                prefix = "- [ ] " if a_s == "□" else "- "
                out.append(prefix + _strip_leading_bullet(b_s))
            elif a_s in ("说明", "判定口径", "备注"):
                flush_table()
                out.append("")
                out.append("**%s**：%s" % (a_s, _md_cell(b_s)))
            else:
                table_buf.append((a_s, b_s))
        # 全空行跳过
    flush_table()
    return "\n".join(out).strip() + "\n"


# ----------------------------- 表格类 sheet 渲染 -----------------------------
def render_table_sheet(ws, title, level=2):
    """渲染 附表X / 超限测点清单 这类"表头块 + 数据表 + 表尾分析"的 sheet。"""
    grid = []
    for r in range(1, ws.max_row + 1):
        grid.append([ws.cell(r, c).value for c in range(1, ws.max_column + 1)])

    counts = [_nonempty_count(row) for row in grid[:8]] or [0]
    max_cnt = max(counts)
    header_idx = next((i for i, c in enumerate(counts) if c == max_cnt), 0)

    out = []
    out.append("#" * level + " " + title)

    # 表头块（表头行之前的单格文本行）
    for i in range(header_idx):
        vals = [str(v).strip() for v in grid[i] if v is not None and str(v).strip() != ""]
        if vals:
            out.append("**" + " ".join(vals) + "**")
    out.append("")

    # 数据表
    header = grid[header_idx]
    ncol = ws.max_column
    out.append("| " + " | ".join(_md_cell(header[c]) for c in range(ncol)) + " |")
    out.append("| " + " | ".join(["---"] * ncol) + " |")
    i = header_idx + 1
    while i < len(grid) and _nonempty_count(grid[i]) >= 2:
        out.append("| " + " | ".join(_md_cell(grid[i][c]) for c in range(ncol)) + " |")
        i += 1

    # 表尾 + 简要分析（单格文本行）
    tail = []
    for j in range(i, len(grid)):
        vals = [str(v).strip() for v in grid[j] if v is not None and str(v).strip() != ""]
        if vals:
            tail.append(vals)
    if tail:
        out.append("")
        for vals in tail:
            line = " ".join(vals)
            if line.strip() == "简要分析及判断性结论":
                out.append("### " + line.strip())
            else:
                out.append(line)
    return "\n".join(out).strip() + "\n"


# ----------------------------- 可复用导出函数 -----------------------------
def export(xlsx, cover_sheet="监测日报", out_cover=None, out_rest=None):
    """把日报 xlsx 导出为两份 Markdown，返回 (封面路径, 附表及清单路径)。"""
    stem = os.path.splitext(os.path.basename(xlsx))[0]
    out_dir = os.path.dirname(os.path.abspath(xlsx))
    out_cover = out_cover or os.path.join(out_dir, stem + "_封面.md")
    out_rest = out_rest or os.path.join(out_dir, stem + "_附表及清单.md")

    wb = openpyxl.load_workbook(xlsx, data_only=True)
    names = wb.sheetnames
    if cover_sheet not in names:
        raise SystemExit("未在 %s 中找到封面 sheet「%s」，实际 sheet：%s" % (xlsx, cover_sheet, names))

    cover_md = render_cover(wb[cover_sheet])
    with open(out_cover, "w", encoding="utf-8") as f:
        f.write(cover_md)

    rest_parts = ["# 监测日报 · 附表与超限测点清单（%s）" % stem, ""]
    for s in names:
        if s == cover_sheet:
            continue
        rest_parts.append(render_table_sheet(wb[s], s, level=2))
        rest_parts.append("")
    rest_md = "\n".join(rest_parts).strip() + "\n"
    with open(out_rest, "w", encoding="utf-8") as f:
        f.write(rest_md)

    print("封面 Markdown :", out_cover)
    print("附表 Markdown :", out_rest)
    print("sheet 分组    : 封面=[%s]｜其余=%s" % (cover_sheet, [s for s in names if s != cover_sheet]))
    return out_cover, out_rest


# ----------------------------- 主流程 -----------------------------
def main():
    ap = argparse.ArgumentParser(description="导出监测日报 xlsx 为两份 Markdown（封面 / 附表及清单）")
    ap.add_argument("--xlsx", required=True, help="日报 xlsx 路径（gen_daily_report_rag.py 产出）")
    ap.add_argument("--cover-sheet", default="监测日报", help="封面 sheet 名（默认 监测日报）")
    ap.add_argument("--out-cover", default=None, help="封面 Markdown 输出路径")
    ap.add_argument("--out-rest", default=None, help="附表及清单 Markdown 输出路径")
    args = ap.parse_args()
    export(args.xlsx, args.cover_sheet, args.out_cover, args.out_rest)


if __name__ == "__main__":
    main()
