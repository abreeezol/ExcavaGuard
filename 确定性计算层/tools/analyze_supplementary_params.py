"""
补充参数剖析（只读）
====================

对《监测参数清单与规范依据》查漏后新增的参数，直接从原始数据提取统计画像：

  1. 土体温度            ← 隧道 observations.csv（监测类参数）
  2. 不排水抗剪强度 Cu   ← 隆起 A3（勘察类参数）
  3. SPT 标准贯入击数 N  ← 隆起 A5（勘察类参数）
  4. 分层测压管水头      ← 隆起 A4（勘察类参数）

约束：
- **只读**：不修改、不覆盖、不删除任何原始文件；
- 结果写入 `derived/补充参数剖析.json`。

运行：
    python -B tools/analyze_supplementary_params.py
"""

from __future__ import annotations

import csv
import json
import statistics as st
from collections import defaultdict
from pathlib import Path

BASE = Path(__file__).resolve().parent.parent
DERIVED = BASE / "derived"
DATA = Path(r"C:\Study\bisai\Hai AI Agent\基坑智守项目相关数据")

TUNNEL = DATA / "多传感器隧道监测数据" / "sensors_public_dataset" / "data"
HEAVE = DATA / "地下室粘土隆起监测"


def _q(vals: list[float], p: float) -> float | None:
    """线性插值分位数。"""
    if not vals:
        return None
    s = sorted(vals)
    if len(s) == 1:
        return s[0]
    k = (len(s) - 1) * p
    lo, hi = int(k), min(int(k) + 1, len(s) - 1)
    return s[lo] + (s[hi] - s[lo]) * (k - lo)


def describe(vals: list[float]) -> dict:
    vals = [v for v in vals if v is not None]
    if not vals:
        return {"n": 0}
    return {
        "n": len(vals),
        "min": round(min(vals), 4),
        "q25": round(_q(vals, 0.25), 4),
        "median": round(st.median(vals), 4),
        "q75": round(_q(vals, 0.75), 4),
        "max": round(max(vals), 4),
        "mean": round(st.fmean(vals), 4),
        "sd": round(st.pstdev(vals), 4) if len(vals) > 1 else 0.0,
    }


# ---------------------------------------------------------------------------
# 1. 土体温度
# ---------------------------------------------------------------------------
def analyze_temperature() -> dict:
    rows = []
    with (TUNNEL / "observations.csv").open(encoding="utf-8") as f:
        for r in csv.DictReader(f):
            try:
                t = float(r["temperature"])
            except (TypeError, ValueError):
                continue
            rows.append((r["timestamp"][:7], r["sensor_type"], t))

    by_type: dict[str, list[float]] = defaultdict(list)
    by_month: dict[str, list[float]] = defaultdict(list)
    for month, stype, t in rows:
        by_type[stype].append(t)
        by_month[month].append(t)

    months = sorted(by_month)
    return {
        "source": "多传感器隧道监测数据/sensors_public_dataset/data/observations.csv",
        "field": "temperature",
        "unit": "℃",
        "note": "时间戳为脱敏后的相对日历，月序仅表示序列位置，非真实月份",
        "overall": describe([t for _, _, t in rows]),
        "by_sensor_type": {k: describe(v) for k, v in sorted(by_type.items())},
        "monthly_mean": {m: round(st.fmean(by_month[m]), 3) for m in months},
        "monthly_span": round(max(st.fmean(by_month[m]) for m in months)
                              - min(st.fmean(by_month[m]) for m in months), 3),
    }


# ---------------------------------------------------------------------------
# 2. 不排水抗剪强度 Cu
# ---------------------------------------------------------------------------
def analyze_cu() -> dict:
    import openpyxl

    wb = openpyxl.load_workbook(HEAVE / "A3 - Undrained shear strength data and calculations.xlsx",
                                data_only=True)
    ws = wb["Undrained Shear Strength"]

    groups = [
        (0, "开挖区中心钻孔"), (2, "开挖区内钻孔"), (4, "开挖区外钻孔"),
        (6, "1962 年勘察"), (8, "设计线 Cu = 75 + 6z"),
    ]
    out: dict[str, dict] = {}
    for col, name in groups:
        pairs = []
        for row in ws.iter_rows(min_row=3, values_only=True):
            if col + 1 >= len(row):
                continue
            cu, rl = row[col], row[col + 1]
            if isinstance(cu, (int, float)) and isinstance(rl, (int, float)):
                pairs.append((float(cu), float(rl)))
        out[name] = {
            "n": len(pairs),
            "cu_kPa": describe([p[0] for p in pairs]),
            "reduced_level_mOD": describe([p[1] for p in pairs]),
        }

    ws2 = wb["Statistical test"]
    stat_rows = []
    for row in ws2.iter_rows(min_row=1, max_row=8, values_only=True):
        stat_rows.append([c for c in row if c is not None])
    wb.close()

    return {
        "source": "地下室粘土隆起监测/A3 - Undrained shear strength data and calculations.xlsx",
        "sheet": "Undrained Shear Strength",
        "unit": "kPa",
        "groups": out,
        "published_stats": {
            "under_excavation_centre": {"mean": 123.33, "sd": 39.92, "n": 37},
            "not_under_excavation_1962SI": {"mean": 131.25, "sd": 50.51, "n": 13},
            "note": "摘自同文件 Statistical test 表；t 检验 p≈0.307，两组均值无显著差异",
        },
        "design_line": "Cu = 75 + 6z（kPa，z 为深度）",
    }


# ---------------------------------------------------------------------------
# 3. SPT 标准贯入击数
# ---------------------------------------------------------------------------
def analyze_spt() -> dict:
    import openpyxl

    wb = openpyxl.load_workbook(
        HEAVE / "A5 - Standard penetration test data and calculations.xlsx", data_only=True)
    ws = wb["SPTs"]
    recs = []
    for row in ws.iter_rows(min_row=4, values_only=True):
        if len(row) < 5:
            continue
        bh, gl, top_lc, depth, spt = row[0], row[1], row[2], row[3], row[4]
        if not isinstance(spt, (int, float)):
            continue
        recs.append({
            "bh": str(bh).strip() if bh is not None else "",
            "gl_mOD": float(gl) if isinstance(gl, (int, float)) else None,
            "top_lc_mbgl": float(top_lc) if isinstance(top_lc, (int, float)) else None,
            "depth_m": float(depth) if isinstance(depth, (int, float)) else None,
            "spt_n": float(spt),
        })
    wb.close()

    by_bh: dict[str, list[float]] = defaultdict(list)
    for r in recs:
        if r["bh"]:
            by_bh[r["bh"]].append(r["spt_n"])

    return {
        "source": "地下室粘土隆起监测/A5 - Standard penetration test data and calculations.xlsx",
        "sheet": "SPTs",
        "unit": "击",
        "overall": describe([r["spt_n"] for r in recs]),
        "borehole_count": len(by_bh),
        "by_borehole": {k: describe(v) for k, v in sorted(by_bh.items())},
        "depth_m": describe([r["depth_m"] for r in recs if r["depth_m"] is not None]),
        "gl_mOD": describe([r["gl_mOD"] for r in recs if r["gl_mOD"] is not None]),
        "note": "GL = 地表高程；Top of LC = 伦敦黏土层面埋深",
    }


# ---------------------------------------------------------------------------
# 4. 分层测压管水头
# ---------------------------------------------------------------------------
def analyze_piezometric() -> dict:
    import openpyxl

    wb = openpyxl.load_workbook(
        HEAVE / "A4 - Piezometric profile with interpreted trendline.xlsx", data_only=True)
    ws = wb["Piezometric data"]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    aquifers = []
    for col in range(0, min(8, len(rows[0]) if rows else 0), 3):
        name = rows[0][col]
        if not name:
            continue
        heads, elevs = [], []
        for r in rows[2:]:
            if col + 1 >= len(r):
                continue
            h, e = r[col], r[col + 1]
            if isinstance(h, (int, float)):
                heads.append(float(h))
            if isinstance(e, (int, float)):
                elevs.append(float(e))
        aquifers.append({
            "aquifer": str(name),
            "head_m": describe(heads),
            "elevation_mOD": describe(elevs),
        })

    return {
        "source": "地下室粘土隆起监测/A4 - Piezometric profile with interpreted trendline.xlsx",
        "sheet": "Piezometric data",
        "unit": "m / mOD",
        "aquifers": aquifers,
        "note": "Head = 测压水头(m)，Elevation = 测点高程(mOD)；用于渗透稳定与降水设计",
    }


def main() -> int:
    print("=" * 74)
    print("补充参数剖析（只读，原始文件不修改）")
    print("=" * 74)

    result = {
        "schema": "excavaguard.supplementary_params/v1",
        "readonly": True,
        "purpose": "对《监测参数清单与规范依据》查漏后新增参数，直接从原始数据提取统计画像",
        "params": {
            "soil_temperature": analyze_temperature(),
            "undrained_shear_strength": analyze_cu(),
            "spt_n_value": analyze_spt(),
            "piezometric_head": analyze_piezometric(),
        },
    }

    t = result["params"]["soil_temperature"]["overall"]
    print(f"\n[1] 土体温度       n={t['n']}  范围 {t['min']}~{t['max']} ℃  均值 {t['mean']}  月均跨度 {result['params']['soil_temperature']['monthly_span']} ℃")

    cu = result["params"]["undrained_shear_strength"]["groups"]
    print("[2] 不排水抗剪强度 Cu（kPa）")
    for k, v in cu.items():
        if v["n"]:
            print(f"      {k:<22} n={v['n']:>3}  均值 {v['cu_kPa']['mean']:>8}  SD {v['cu_kPa']['sd']:>7}  范围 {v['cu_kPa']['min']}~{v['cu_kPa']['max']}")

    spt = result["params"]["spt_n_value"]
    print(f"[3] SPT 击数       n={spt['overall']['n']}  范围 {spt['overall']['min']}~{spt['overall']['max']}  均值 {spt['overall']['mean']}  钻孔 {spt['borehole_count']} 个")

    print("[4] 分层测压管水头")
    for a in result["params"]["piezometric_head"]["aquifers"]:
        print(f"      {a['aquifer']:<16} 水头 {a['head_m'].get('min')}~{a['head_m'].get('max')} m   高程 {a['elevation_mOD'].get('min')}~{a['elevation_mOD'].get('max')} mOD")

    DERIVED.mkdir(parents=True, exist_ok=True)
    out = DERIVED / "补充参数剖析.json"
    out.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"\n结果已写入：{out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
